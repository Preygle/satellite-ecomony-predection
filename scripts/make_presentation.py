"""Build the 10-slide review presentation.

    python scripts/make_presentation.py

Every figure comes from the project's own outputs. Slide images are the
dataset renders in ``dataset_viewer/map``.
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
IMG = ROOT / "dataset_viewer" / "map"
OUT = ROOT / "docs" / "Review_Presentation.pptx"

# Light theme. Contrast against the white page is the constraint here: the
# blue and amber both had to be darkened from their dark-theme values, which
# were tuned for legibility against near-black and read as washed out on white.
BG        = RGBColor(0xFF, 0xFF, 0xFF)   # page
PANEL     = RGBColor(0xF1, 0xF5, 0xFA)   # card fill
TEXT      = RGBColor(0x10, 0x18, 0x26)   # body text
DIM       = RGBColor(0x5A, 0x64, 0x74)   # secondary text
ACCENT    = RGBColor(0x1B, 0x5F, 0xAF)   # headings, emphasis
WARN      = RGBColor(0xB2, 0x6A, 0x00)   # cautions, key results
BORDER    = RGBColor(0xD5, 0xDD, 0xE8)   # card outline
CHIP_BG   = RGBColor(0xFF, 0xFF, 0xFF)   # chip fill
CHIP_LINE = RGBColor(0xC3, 0xCE, 0xDC)   # chip outline
GREEN     = RGBColor(0x0E, 0x7A, 0x52)   # the no-account data path
ARROW     = RGBColor(0xB8, 0xC4, 0xD4)   # flow arrows

W, H = Inches(13.333), Inches(7.5)


def blank(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.shapes.add_shape(1, 0, 0, W, H)
    bg.fill.solid(); bg.fill.fore_color.rgb = BG
    bg.line.fill.background(); bg.shadow.inherit = False
    return s


def tb(slide, x, y, w, h, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.paragraphs[0].alignment = align
    return tf


def para(tf, text, size=16, colour=TEXT, bold=False, space=8, first=False, align=None):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.text = text
    p.space_after = Pt(space)
    if align is not None:
        p.alignment = align
    for r in p.runs:
        r.font.size = Pt(size); r.font.color.rgb = colour
        r.font.bold = bold; r.font.name = "Segoe UI"
    return p


def heading(slide, title, kicker=None):
    tf = tb(slide, Inches(0.7), Inches(0.42), Inches(12), Inches(0.9))
    para(tf, title, 30, TEXT, True, 2, first=True)
    if kicker:
        para(tf, kicker, 13.5, ACCENT, False, 0)
    bar = slide.shapes.add_shape(1, Inches(0.7), Inches(1.42), Inches(1.5), Pt(3))
    bar.fill.solid(); bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background(); bar.shadow.inherit = False


def card(slide, x, y, w, h):
    c = slide.shapes.add_shape(1, x, y, w, h)
    c.fill.solid(); c.fill.fore_color.rgb = PANEL
    c.line.color.rgb = BORDER; c.line.width = Pt(0.75)
    c.shadow.inherit = False
    return c


def bullets(tf, items, size=15, gap=9):
    for i, it in enumerate(items):
        if isinstance(it, tuple):
            label, body = it
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(gap)
            r1 = p.add_run(); r1.text = label + "  "
            r1.font.size = Pt(size); r1.font.bold = True
            r1.font.color.rgb = TEXT; r1.font.name = "Segoe UI"
            r2 = p.add_run(); r2.text = body
            r2.font.size = Pt(size); r2.font.color.rgb = DIM; r2.font.name = "Segoe UI"
        else:
            para(tf, it, size, DIM, False, gap, first=(i == 0))


def picture(slide, name, x, y, w):
    p = IMG / name
    if p.exists():
        return slide.shapes.add_picture(str(p), x, y, width=w)
    return None


def paper_card(slide, x, y, w, h, title, meta, method, result):
    """One reviewed paper: its title, when it was published, and what it found."""
    c = card(slide, x, y, w, h)
    c.line.color.rgb = CHIP_LINE

    pad = Inches(0.2)
    inner = w - Inches(0.4)

    tf = tb(slide, x + pad, y + Inches(0.12), inner, Inches(0.62))
    para(tf, title, 12.5, TEXT, True, 2, first=True)

    tf = tb(slide, x + pad, y + Inches(0.72), inner, Inches(0.26))
    para(tf, meta, 10.5, ACCENT, True, 0, first=True)

    tf = tb(slide, x + pad, y + Inches(1.0), inner, Inches(0.9))
    p = tf.paragraphs[0]
    r1 = p.add_run(); r1.text = "Method  "
    r1.font.size = Pt(10.5); r1.font.bold = True
    r1.font.color.rgb = DIM; r1.font.name = "Segoe UI"
    r2 = p.add_run(); r2.text = method
    r2.font.size = Pt(10.5); r2.font.color.rgb = DIM; r2.font.name = "Segoe UI"

    p2 = tf.add_paragraph(); p2.space_before = Pt(4)
    r3 = p2.add_run(); r3.text = "Result  "
    r3.font.size = Pt(10.5); r3.font.bold = True
    r3.font.color.rgb = WARN; r3.font.name = "Segoe UI"
    r4 = p2.add_run(); r4.text = result
    r4.font.size = Pt(10.5); r4.font.color.rgb = TEXT; r4.font.name = "Segoe UI"


def paper_grid(slide, papers):
    """Two-by-two grid of paper cards."""
    xs = [Inches(0.7), Inches(6.9)]
    ys = [Inches(1.72), Inches(4.28)]
    for i, pr in enumerate(papers):
        paper_card(slide, xs[i % 2], ys[i // 2], Inches(5.7), Inches(2.4), *pr)


# ---------------------------------------------------------------- slides ---
def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H

    # 1 — Title -------------------------------------------------------------
    s = blank(prs)
    pic = picture(s, "truecolour.png", Inches(7.6), Inches(0), Inches(5.8))
    if pic:
        pic.left, pic.top = Emu(int(W) - pic.width), Emu(int((int(H) - pic.height) / 2))
    tf = tb(s, Inches(0.8), Inches(1.5), Inches(6.6), Inches(4))
    para(tf, "Satellite-Based Urban Growth and", 34, TEXT, True, 0, first=True)
    para(tf, "Economic Activity Intelligence System", 34, TEXT, True, 10)
    para(tf, "Detecting built-but-inactive development from open satellite data", 16, ACCENT, False, 26)
    para(tf, "Case study:  Varanasi, Uttar Pradesh", 15, DIM, False, 30)
    para(tf, "Team", 12, ACCENT, True, 4)
    para(tf, "[Name 1]  —  [Registration Number]", 15, TEXT, False, 3)
    para(tf, "[Name 2]  —  [Registration Number]", 15, TEXT, False, 3)
    para(tf, "[Name 3]  —  [Registration Number]", 15, TEXT, False, 14)
    para(tf, "Guide:  [Guide Name]     ·     BCSE497J Project I", 13, DIM, False, 2)
    para(tf, "School of Computer Science and Engineering  ·  Fall 2026–27", 13, DIM, False, 0)

    # 2 — Literature: nighttime light ---------------------------------------
    s = blank(prs)
    heading(s, "Literature Review  ·  1 of 3",
            "Nighttime light as a proxy for human activity  —  20 papers reviewed, 19 from 2020 onward")
    paper_grid(s, [
        ("A global annual simulated VIIRS nighttime light dataset from 1992 to 2023",
         "Chen, X. et al.  ·  2024  ·  Scientific Data",
         "U-Net convolutional network reframed as image super-resolution, using Landsat NDVI as a second input.",
         "R² 0.617 per pixel rising to 0.964 per country. Error concentrates at the urban fringe."),
        ("An extended time series (2000–2018) of global NPP-VIIRS-like nighttime light data",
         "Chen, Z. et al.  ·  2021  ·  Earth System Science Data",
         "Auto-encoder neural network combined with a vegetation index to break sensor saturation.",
         "R² 0.87 per pixel, 0.95 per city, validated on 150,000 random pixels."),
        ("Cross-Sensor Nighttime Lights Image Calibration with Residual U-Net",
         "Nechaev, D. et al.  ·  2021  ·  Remote Sensing",
         "Residual U-Net converting the newer sensor back to match the older one — the easy direction.",
         "R² 0.94–0.99. Confirms the reverse direction is what is genuinely hard."),
        ("A Prolonged Artificial Nighttime-light Dataset of China (1984–2020)",
         "Zhang, L. et al.  ·  2024  ·  Scientific Data",
         "Convolutional Long Short-Term Memory network, learning how light changes over time.",
         "RMSE 0.73, R² 0.95. Scores higher because it never crosses between two sensors."),
    ])

    # 3 — Literature: built-up extent and vacancy ---------------------------
    s = blank(prs)
    heading(s, "Literature Review  ·  2 of 3",
            "Measuring built-up extent, and inferring whether it is used")
    paper_grid(s, [
        ("Using Landsat and nighttime lights for supervised pixel-based image classification",
         "Goldblatt, R. et al.  ·  2018  ·  Remote Sensing of Environment",
         "Nighttime light used to label training data automatically, avoiding hand-drawn training areas.",
         "80.8% balanced accuracy for India. The methodological ancestor of this project."),
        ("Continental-Scale Building Detection from High Resolution Satellite Imagery",
         "Sirko, W. et al.  ·  2021  ·  Google Research",
         "U-Net on 50 cm imagery, trained on 100,000 images with 1.75 million labelled buildings.",
         "516 million footprints. Deeper encoders gave no gain — technique beat architecture."),
        ("Dynamic World, Near real-time global 10 m land use land cover mapping",
         "Brown, C. F. et al.  ·  2022  ·  Scientific Data",
         "Deep learning on Sentinel-2, publishing a class probability rather than one fixed label.",
         "First near-real-time global land cover. Supplies the water mask our heat layer needs."),
        ("Inferring ghost cities on the globe in newly developed urban areas",
         "Zhang, Y., Tu, T. & Long, Y.  ·  2024  ·  Cities",
         "Urban vitality from road density, points of interest and population, across 8,841 cities.",
         "New areas show only 7.69% of the vitality of old ones; worst 5% called ghost cities."),
    ])

    # 4 — Literature: prediction, environment, and the gap ------------------
    s = blank(prs)
    heading(s, "Literature Review  ·  3 of 3",
            "Predicting expansion, measuring its consequences  —  and the gap we found")
    xs = [Inches(0.7), Inches(6.9)]
    papers = [
        ("Understanding the drivers of sustainable land expansion using the PLUS model",
         "Liang, X. et al.  ·  2021  ·  Computers, Environment and Urban Systems",
         "Random forest to find drivers, then a cellular automaton growing realistic patches.",
         "Figure of Merit 0.2642 — the benchmark our own 0.0679 is measured against."),
        ("An efficient built-up land expansion model using a modified U-Net",
         "Shojaei, H. et al.  ·  2022  ·  International Journal of Digital Earth",
         "Modified U-Net labelling every pixel built-up or not, from slope and distance drivers.",
         "AUC 0.87 against 0.82 for random forest — a 0.05 gain, at the cost of a GPU."),
        ("Spatio-temporal analysis of urban expansion using Google Earth Engine",
         "Zhang, A. et al.  ·  2025  ·  Scientific Reports",
         "Random forest in Earth Engine, then a cellular automaton with a neural network, to 2030.",
         "Accuracy above 92%. The closest recent precedent for our whole workflow."),
        ("Google Earth Engine Open-Source Code for Landsat Land Surface Temperature",
         "Ermida, S. L. et al.  ·  2020  ·  Remote Sensing",
         "Open code computing land surface temperature from Landsat entirely inside Earth Engine.",
         "RMSE 1.0–1.3 K against ground sensors. The reference method behind our heat layer."),
    ]
    ys = [Inches(1.72), Inches(3.92)]
    for i, pr in enumerate(papers):
        paper_card(s, xs[i % 2], ys[i // 2], Inches(5.7), Inches(2.05), *pr)

    gap = card(s, Inches(0.7), Inches(6.15), Inches(11.9), Inches(1.0))
    gap.line.color.rgb = WARN
    tf = tb(s, Inches(1.0), Inches(6.26), Inches(11.3), Inches(0.85))
    para(tf, "The gap:  no study learns the activity level normal for a given level of "
             "development within a city, then uses it to separate development that has "
             "stalled from development still filling up.", 14, TEXT, True, 4, first=True)
    para(tf, "On our own data that distinction is the difference between 7.24 km² and "
             "0.35 km² — a twentyfold overstatement without it.", 12, WARN, False, 0)


    # 5 — Methodology -------------------------------------------------------
    s = blank(prs)
    heading(s, "Methodology", "Six steps, each a separate testable module")
    steps = [
        ("1", "Find the city",
         "A city name is turned into coordinates, and a circle around that point becomes the study area.",
         "A city has no single official boundary available everywhere, and municipal limits usually exclude the fastest-growing fringe — which is exactly where our findings are."),
        ("2", "Download the data",
         "GHSL and OpenStreetMap come straight over HTTP; the satellite layers come through Google Earth Engine.",
         "Earth Engine clips each layer on its own servers, so a few megabytes arrive instead of whole scenes. The GHSL path needs no account, so the system always returns a result."),
        ("3", "Put it on one grid",
         "Every dataset is reprojected onto a single 100 m grid in UTM zone 44N, and reported at 500 m.",
         "Built-up area and population are extensive quantities, so they are converted to density before reprojection and rescaled after — otherwise the totals would not survive the warp."),
        ("4", "Compute indicators",
         "Built-up change, vegetation loss, surface temperature, night-light trend, and shop and road density.",
         "Each is normalised per unit of built-up area, because raw values otherwise measure how large a place is rather than how intensively it is used."),
        ("5", "Train the model",
         "Logistic Regression learns which cells became built-up, from seven drivers; a cellular automaton places the growth.",
         "The automaton makes each cell look at its neighbours, so predicted development appears in connected patches rather than scattered single cells — which is how cities actually grow."),
        ("6", "Predict and display",
         "Future demand is extrapolated from the city's own past growth rate, then allocated to the most suitable cells.",
         "The +5 and +10 year maps are a prediction, not a measurement — they show where growth is likely if current trends continue, and the dashboard says so."),
    ]
    y = Inches(1.78)
    for n, title, body, why in steps:
        card(s, Inches(0.7), y, Inches(11.9), Inches(0.79))
        num = s.shapes.add_shape(9, Inches(0.88), y + Inches(0.19), Inches(0.42), Inches(0.42))
        num.fill.solid(); num.fill.fore_color.rgb = ACCENT
        num.line.fill.background(); num.shadow.inherit = False
        ntf = num.text_frame; ntf.word_wrap = False
        np_ = ntf.paragraphs[0]; np_.text = n; np_.alignment = PP_ALIGN.CENTER
        for r in np_.runs:
            r.font.size = Pt(14); r.font.bold = True
            r.font.color.rgb = BG; r.font.name = "Segoe UI"
        tf = tb(s, Inches(1.45), y + Inches(0.08), Inches(2.35), Inches(0.62))
        para(tf, title, 14, TEXT, True, 0, first=True)
        tf = tb(s, Inches(3.8), y + Inches(0.06), Inches(8.6), Inches(0.72))
        para(tf, body, 11.5, TEXT, False, 3, first=True)
        para(tf, why, 10.5, DIM, False, 0)
        y = Emu(int(y) + int(Inches(0.84)))

    tf = tb(s, Inches(0.7), Inches(6.65), Inches(11.9), Inches(0.6))
    para(tf, "Validation:  train on 2010→2015, test on the unseen 2015→2020, and compare "
             "against random allocation. Figure of Merit 0.0679 versus a random baseline "
             "of 0.0055 — 12.3× better than chance.", 13, WARN, False, 0, first=True)

    # 6 — Workflow ----------------------------------------------------------
    s = blank(prs)
    heading(s, "Workflow", "Seven pipeline stages, run in order")
    stages = [("builtup", "GHS-BUILT-S\nGHS-POP"), ("osm", "OpenStreetMap\nOverpass API"),
              ("gee", "VIIRS · Sentinel-2\nLandsat · Dynamic World"), ("vegetation", "Green cover\nchange"),
              ("thermal", "Heat island\nintensity"), ("ghost", "Activity index\n+ typology"),
              ("export", "500 m grid\n+ summary")]
    x = Inches(0.7)
    for i, (name, body) in enumerate(stages):
        c = card(s, x, Inches(2.1), Inches(1.55), Inches(1.9))
        c.line.color.rgb = ACCENT if i in (2, 5) else BORDER
        tf = tb(s, x + Inches(0.1), Inches(2.25), Inches(1.35), Inches(0.4), PP_ALIGN.CENTER)
        para(tf, f"{i+1}. {name}", 12.5, ACCENT, True, 6, first=True, align=PP_ALIGN.CENTER)
        para(tf, body, 11, DIM, False, 0, align=PP_ALIGN.CENTER)
        x = Emu(int(x) + int(Inches(1.72)))

    card(s, Inches(0.7), Inches(4.35), Inches(5.9), Inches(2.3))
    tf = tb(s, Inches(1.0), Inches(4.55), Inches(5.3), Inches(2.0))
    para(tf, "A failed stage does not stop the run", 16, ACCENT, True, 10, first=True)
    para(tf, "Public data sources rate-limit and time out, and Earth Engine needs an "
             "account. Any stage that cannot run records why, and the pipeline continues. "
             "The reason appears in the output and on the dashboard, so a missing layer is "
             "always visible rather than silently read as zero.", 13, DIM, False, 0)

    card(s, Inches(6.9), Inches(4.35), Inches(5.7), Inches(2.3))
    tf = tb(s, Inches(7.2), Inches(4.55), Inches(5.1), Inches(2.0))
    para(tf, "Prediction runs separately", 16, ACCENT, True, 10, first=True)
    para(tf, "Build seven drivers  →  train on an early period  →  test on a later unseen "
             "period  →  compare against random  →  project +5 and +10 years  →  allocate "
             "with a cellular automaton so growth appears in connected patches.",
         13, DIM, False, 0)

    # 7 — Architecture ------------------------------------------------------
    s = blank(prs)
    heading(s, "System Architecture", "Four layers · data flows upward")
    def band(y, h, label, colour):
        c = card(s, Inches(0.7), y, Inches(8.9), h)
        c.line.color.rgb = colour
        tag = s.shapes.add_shape(5, Inches(0.7), y, Inches(0.16), h)
        tag.fill.solid(); tag.fill.fore_color.rgb = colour
        tag.line.fill.background(); tag.shadow.inherit = False
        tf = tb(s, Inches(1.05), y + Inches(0.11), Inches(3.0), Inches(0.32))
        para(tf, label, 11.5, colour, True, 0, first=True)
        return c

    def chip(x, y, w, text, colour=DIM, h=Inches(0.36)):
        c = s.shapes.add_shape(5, x, y, w, h)
        c.fill.solid(); c.fill.fore_color.rgb = CHIP_BG
        c.line.color.rgb = CHIP_LINE; c.line.width = Pt(0.75)
        c.shadow.inherit = False
        tf = c.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.text = text; p.alignment = PP_ALIGN.CENTER
        for r in p.runs:
            r.font.size = Pt(9.5); r.font.color.rgb = colour; r.font.name = "Segoe UI"
        return c

    def up_arrow(y):
        a = s.shapes.add_shape(34, Inches(4.85), y, Inches(0.55), Inches(0.26))
        a.fill.solid(); a.fill.fore_color.rgb = ARROW
        a.line.fill.background(); a.shadow.inherit = False

    # Presentation ----------------------------------------------------------
    band(Inches(1.72), Inches(0.68), "PRESENTATION", ACCENT)
    for i, t_ in enumerate(["Interactive map", "Growth typology", "Ghost zones",
                            "Environment", "Grid export"]):
        chip(Inches(4.2 + i * 1.08), Inches(1.87), Inches(1.02), t_, TEXT)

    up_arrow(Inches(2.48))

    # Analysis --------------------------------------------------------------
    band(Inches(2.82), Inches(1.22), "ANALYSIS", TEXT)
    mods = ["builtup.py", "nightlights.py", "vegetation.py",
            "thermal.py", "ghost.py", "growth_model.py"]
    for i, m in enumerate(mods):
        chip(Inches(4.2 + (i % 3) * 1.78), Inches(2.98 + (i // 3) * 0.5),
             Inches(1.68), m, TEXT)
    tf = tb(s, Inches(1.05), Inches(3.32), Inches(3.0), Inches(0.6))
    para(tf, "Never downloads —\nso it is testable offline", 9.5, DIM, False, 0, first=True)

    up_arrow(Inches(4.14))

    # Acquisition — two paths ------------------------------------------------
    band(Inches(4.48), Inches(1.42), "ACQUISITION", TEXT)
    tf = tb(s, Inches(4.2), Inches(4.56), Inches(2.6), Inches(0.3))
    para(tf, "OPEN PATH — no account", 9.5, GREEN, True, 0, first=True)
    for i, t_ in enumerate(["GHS-BUILT-S", "GHS-POP", "OpenStreetMap"]):
        chip(Inches(4.2), Inches(4.86 + i * 0.32), Inches(2.3), t_, TEXT, Inches(0.28))
    tf = tb(s, Inches(6.85), Inches(4.56), Inches(2.6), Inches(0.3))
    para(tf, "EARTH ENGINE PATH — one login", 9.5, WARN, True, 0, first=True)
    for i, t_ in enumerate(["VIIRS nighttime lights", "Sentinel-2", "Landsat 8 / 9",
                            "Dynamic World", "Open Buildings"]):
        chip(Inches(6.85), Inches(4.86 + i * 0.20), Inches(2.55), t_, TEXT, Inches(0.18))

    up_arrow(Inches(5.99))

    # Configuration ----------------------------------------------------------
    band(Inches(6.33), Inches(0.68), "CONFIGURATION", DIM)
    tf = tb(s, Inches(4.2), Inches(6.44), Inches(5.2), Inches(0.5))
    para(tf, "config/varanasi.yaml  —  city, area, grid, epochs, thresholds, sources",
         10.5, TEXT, False, 0, first=True)

    # Side commentary ---------------------------------------------------------
    notes = [
        ("Two paths, on purpose", ACCENT,
         "Most projects of this kind stall at “first, get an Earth Engine account”. "
         "Splitting the sources means the core result exists before anyone logs in."),
        ("Layers only talk downward", TEXT,
         "Acquisition never analyses; analysis never downloads. That is why the analysis "
         "can be tested on synthetic data with no network — the suite runs in seconds."),
        ("A new city is a new file", WARN,
         "Nothing below the configuration layer is hard-coded. Pointing the system at "
         "another city changes a centre point and a radius, not a line of code."),
    ]
    y = Inches(1.72)
    for title, colour, body in notes:
        card(s, Inches(9.8), y, Inches(2.8), Inches(1.62))
        tf = tb(s, Inches(10.0), y + Inches(0.13), Inches(2.45), Inches(1.4))
        para(tf, title, 11.5, colour, True, 5, first=True)
        para(tf, body, 9.5, DIM, False, 0)
        y = Emu(int(y) + int(Inches(1.78)))

    # 8 — Datasets ----------------------------------------------------------
    s = blank(prs)
    heading(s, "Datasets Acquired", "Official identifiers, exactly as published")
    data = [
        ("GHS-BUILT-S R2023A", "EC Joint Research Centre", "100 m", "2020 epoch", "Built-up surface"),
        ("GHS-POP R2023A", "EC Joint Research Centre", "100 m", "2020 epoch", "Population"),
        ("NOAA/VIIRS/DNB/ANNUAL_V21 / _V22", "NOAA", "463 m", "2013–2024", "Nighttime lights"),
        ("COPERNICUS/S2_SR_HARMONIZED", "ESA / Copernicus", "10 m", "Oct 2024 – Mar 2025", "True colour, NDVI, NDBI"),
        ("LANDSAT/LC08 & LC09/C02/T1_L2", "USGS", "30 m", "Mar – May 2024", "Land surface temperature"),
        ("GOOGLE/DYNAMICWORLD/V1", "Google", "10 m", "Oct 2024 – Mar 2025", "Land cover"),
        ("GOOGLE/Research/open-buildings-temporal/v1", "Google", "4 m", "2023", "Building height"),
        ("OpenStreetMap · Overpass API", "OSM contributors", "vector", "Current", "Shops, offices, roads"),
    ]
    hdr = ["Official dataset", "Provider", "Res.", "Imagery dates", "Gives us"]
    xs2 = [Inches(0.8), Inches(5.0), Inches(7.3), Inches(8.5), Inches(10.6)]
    ws = [Inches(4.1), Inches(2.2), Inches(1.1), Inches(2.0), Inches(2.1)]
    card(s, Inches(0.7), Inches(1.75), Inches(11.9), Inches(4.5))
    tf = tb(s, Inches(0.8), Inches(1.85), Inches(11.7), Inches(0.35))
    para(tf, "", 8, DIM, False, 0, first=True)
    y = Inches(1.9)
    for i, x, w in zip(range(5), xs2, ws):
        t2 = tb(s, x, y, w, Inches(0.35))
        para(t2, hdr[i], 11.5, ACCENT, True, 0, first=True)
    y = Emu(int(y) + int(Inches(0.42)))
    for row in data:
        for i, (x, w) in enumerate(zip(xs2, ws)):
            t2 = tb(s, x, y, w, Inches(0.42))
            para(t2, row[i], 10.5 if i == 0 else 10.5,
                 TEXT if i == 0 else DIM, i == 0, 0, first=True)
        y = Emu(int(y) + int(Inches(0.5)))

    tf = tb(s, Inches(0.7), Inches(6.45), Inches(11.9), Inches(0.7))
    para(tf, "Acquisition dates are read back from the image collections themselves — the "
             "Sentinel-2 layers are a median of 116 scenes across 29 separate days — not "
             "assumed from a filename.", 12.5, WARN, False, 0, first=True)

    # 9 — Pipeline in pictures ----------------------------------------------
    s = blank(prs)
    heading(s, "The Complete Pipeline", "Real output at every stage — nothing here is a sample image")
    shots = [("truecolour.png", "1 · Satellite image", "Sentinel-2 true colour"),
             ("builtup.png", "2 · Built-up surface", "GHS-BUILT-S, 2020"),
             ("nightlights.png", "3 · Activity", "VIIRS nighttime lights"),
             ("lst.png", "4 · Heat island", "Landsat surface temperature")]
    x = Inches(0.7)
    for name, title, sub in shots:
        card(s, x, Inches(1.8), Inches(2.85), Inches(3.5))
        pic = picture(s, name, x + Inches(0.15), Inches(1.95), Inches(2.55))
        if pic and pic.height > Inches(2.4):
            ratio = int(Inches(2.4)) / pic.height
            pic.height = Inches(2.4); pic.width = Emu(int(pic.width * ratio))
            pic.left = Emu(int(x) + int((int(Inches(2.85)) - pic.width) / 2))
        tf = tb(s, x + Inches(0.15), Inches(4.6), Inches(2.55), Inches(0.7))
        para(tf, title, 13, ACCENT, True, 3, first=True)
        para(tf, sub, 11, DIM, False, 0)
        x = Emu(int(x) + int(Inches(3.0)))

    card(s, Inches(0.7), Inches(5.5), Inches(11.9), Inches(1.5))
    tf = tb(s, Inches(1.0), Inches(5.68), Inches(11.3), Inches(1.2))
    para(tf, "→  combined into one activity index  →  compared against the level normal for "
             "that much development in this city  →  six-class growth typology", 14, TEXT, True, 8, first=True)
    para(tf, "Result:  0.35 km² of genuinely stalled development across 15 zones, all "
             "peripheral — consistent with the finding that 48.8% of new development is "
             "detached from the existing city.", 13, DIM, False, 0)

    # 10 — Conclusion -------------------------------------------------------
    s = blank(prs)
    heading(s, "Conclusion", "What the system delivers, and what it does not claim")
    card(s, Inches(0.7), Inches(1.75), Inches(5.9), Inches(2.6))
    tf = tb(s, Inches(1.0), Inches(1.95), Inches(5.3), Inches(2.2))
    para(tf, "Key findings for Varanasi", 16, ACCENT, True, 10, first=True)
    bullets(tf, [
        ("Built-up +23.8%", "against population +10.3%, 2010–2020"),
        ("2.3× more land", "consumed per resident"),
        ("48.8% leapfrog", "new development detached from the city"),
        ("0.35 km² stalled", "across 15 zones, all peripheral"),
        ("+31.21 km²", "projected urban extent by 2030"),
    ], 13, 8)

    card(s, Inches(6.9), Inches(1.75), Inches(5.7), Inches(2.6))
    tf = tb(s, Inches(7.2), Inches(1.95), Inches(5.1), Inches(2.2))
    para(tf, "Contribution", 16, ACCENT, True, 10, first=True)
    para(tf, "The activity level expected of a place is learned from the city's own data "
             "rather than set as a fixed threshold, so the method transfers to any city "
             "unchanged. Development that has stalled is separated from development still "
             "filling up — a distinction that accounts for 95% of what would otherwise be "
             "reported as a problem.", 13, DIM, False, 0)

    card(s, Inches(0.7), Inches(4.55), Inches(5.9), Inches(2.1))
    tf = tb(s, Inches(1.0), Inches(4.72), Inches(5.3), Inches(1.8))
    para(tf, "What this is not", 16, WARN, True, 10, first=True)
    bullets(tf, [
        "Screening, not a census — the smallest reliable unit is a neighbourhood",
        "No ground validation yet; that is the first Phase 3 task",
        "Figure of Merit 0.0679 carries real signal but is not parcel-accurate",
    ], 12.5, 7)

    card(s, Inches(6.9), Inches(4.55), Inches(5.7), Inches(2.1))
    tf = tb(s, Inches(7.2), Inches(4.72), Inches(5.1), Inches(1.8))
    para(tf, "Next", 16, ACCENT, True, 10, first=True)
    bullets(tf, [
        "Ground validation of the 15 flagged zones",
        "Ward-level reporting on Census 2011 boundaries",
        "City search box replacing the configuration file",
        "Master-plan overlay: planned versus actual development",
    ], 12.5, 7)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    return OUT


if __name__ == "__main__":
    p = build()
    print(f"wrote {p}  ({p.stat().st_size // 1024} KB)")
