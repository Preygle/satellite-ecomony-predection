from __future__ import annotations

import shutil
import sys
from datetime import datetime
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor

ROOT = Path(__file__).resolve().parents[1]

# Dark value -> light replacement. Fills and lines differ where the same dark
# colour served two purposes: 2F3A48 was both an arrow fill and a card outline.
FILL_MAP = {
    "0F1115": "FFFFFF",   # slide background
    "161A20": "F1F5FA",   # card panel
    "11151B": "FFFFFF",   # chip
    "4F9CF9": "1B5FAF",   # accent bar, numbered circle
    "2F3A48": "B8C4D4",   # flow arrow
    "E6EDF3": "101826",   # band tag drawn in text colour
    "9AA4B0": "5A6474",   # band tag drawn in dim colour
    "EDA100": "B26A00",
    "1BAF7A": "0E7A52",
}

LINE_MAP = {
    "262C36": "D5DDE8",   # card outline
    "2F3A48": "C3CEDC",   # paper-card outline
    "303844": "C3CEDC",   # chip outline
    "4F9CF9": "1B5FAF",   # highlighted card outline
    "EDA100": "B26A00",
}

FONT_MAP = {
    "E6EDF3": "101826",   # body text
    "9AA4B0": "5A6474",   # secondary text
    "4F9CF9": "1B5FAF",   # accent
    "EDA100": "B26A00",   # cautions
    "1BAF7A": "0E7A52",   # green path label
    "0F1115": "FFFFFF",   # numbered circles: was dark-on-blue, now white-on-blue
}


def swap(current, mapping):
    if current is None:
        return None
    key = str(current).upper()
    return RGBColor.from_string(mapping[key]) if key in mapping else None


def recolour(path: Path) -> dict:
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    backup = path.with_name(f"{path.stem}.backup-{stamp}{path.suffix}")
    shutil.copy2(path, backup)

    prs = Presentation(str(path))
    counts = {"fill": 0, "line": 0, "font": 0, "shapes": 0}

    for slide in prs.slides:
        for shape in slide.shapes:
            counts["shapes"] += 1

            try:
                if shape.fill.type == 1:                      # solid
                    new = swap(shape.fill.fore_color.rgb, FILL_MAP)
                    if new is not None:
                        shape.fill.fore_color.rgb = new
                        counts["fill"] += 1
            except Exception:
                pass

            try:
                new = swap(shape.line.color.rgb, LINE_MAP)
                if new is not None:
                    shape.line.color.rgb = new
                    counts["line"] += 1
            except Exception:
                pass

            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    for run in para.runs:
                        try:
                            new = swap(run.font.color.rgb, FONT_MAP)
                            if new is not None:
                                run.font.color.rgb = new
                                counts["font"] += 1
                        except Exception:
                            pass

    prs.save(str(path))
    counts["backup"] = backup.name
    return counts


def main() -> int:
    if len(sys.argv) < 2:
        print("usage: python scripts/recolour_deck.py <deck.pptx>   (edits in place, backup first)")
        return 2
    p = Path(sys.argv[1])
    if not p.is_absolute():
        p = ROOT / p
    if not p.exists():
        print(f"not found: {p}")
        return 2
    c = recolour(p)
    print(f"recoloured {p.name}")
    print(f"  shapes visited : {c['shapes']}")
    print(f"  fills changed  : {c['fill']}")
    print(f"  lines changed  : {c['line']}")
    print(f"  fonts changed  : {c['font']}")
    print(f"  backup written : {c['backup']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
