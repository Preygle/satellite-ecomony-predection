from __future__ import annotations

import json
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from make_presentation import (  # noqa: E402
    ACCENT, BORDER, DIM, GREEN, TEXT, WARN, H, W, blank, bullets, card, heading, para, tb,
)
from make_presentation_r3 import chip, footer, grid_table  # noqa: E402

DIAG = ROOT / "docs" / "diagrams" / "presentation"
FIG = ROOT / "docs" / "figures"
OUT = ROOT / "docs" / "Review3_Panel_Presentation.pptx"
RESULTS = ROOT / "outputs" / "review3_results.json"

R = json.loads(RESULTS.read_text(encoding="utf-8")) if RESULTS.exists() else {}


def g(path: str, default):
    """Dotted lookup into the results record, with the value the deck falls back to."""
    node = R
    for part in path.split("."):
        if not isinstance(node, dict) or part not in node:
            return default
        node = node[part]
    return node


# --------------------------------------------------------------- helpers ---

def place(slide, path: Path, x, y, w, h, middle=False):
    """Fit an image inside the box, centred across it; top-aligned unless `middle`."""
    if not path.exists():
        print(f"  missing image: {path.relative_to(ROOT)}")
        return None
    iw, ih = Image.open(path).size
    scale = min(w / iw, h / ih)
    pw, ph = int(iw * scale), int(ih * scale)
    top = y + int((h - ph) / 2) if middle else y
    return slide.shapes.add_picture(str(path), x + int((w - pw) / 2), top, width=pw, height=ph)


def notes(slide, text: str):
    slide.notes_slide.notes_text_frame.text = text


def tidy(prs):
    """Proper unit symbols on the slides (the source strings stay ASCII-friendly)."""
    swaps = (("km2", "km²"), ("R2 ", "R² "), ("degrees C", "°C"))
    for slide in prs.slides:
        for shape in slide.shapes:
            frames = [shape.text_frame] if shape.has_text_frame else []
            if shape.has_table:
                frames += [cell.text_frame for row in shape.table.rows for cell in row.cells]
            for tf in frames:
                for p in tf.paragraphs:
                    for r in p.runs:
                        for a, b in swaps:
                            r.text = r.text.replace(a, b)


def diagram_slide(prs, title, claim, name, talk):
    """One diagram, as large as the slide allows; the claim sits under the title."""
    s = blank(prs)
    heading(s, title, claim)
    place(s, DIAG / name, Inches(0.5), Inches(1.65), Inches(12.33), Inches(5.55), middle=True)
    notes(s, talk)
    return s


# ---------------------------------------------------------------- slides ---

def build():
    b = R.get("builtup", {})
    built = b.get("built_surface_km2", {})
    proj = R.get("projections", {})
    p25 = proj.get("2025", {}).get("demand_km2", 14.89)
    p30 = proj.get("2030", {}).get("demand_km2", 31.21)
    rf = g("growth_models.random_forest", {})
    lr = g("growth_models.logistic_regression", {})

    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H

    # 1 - title ------------------------------------------------------------
    s = blank(prs)
    tf = tb(s, Inches(0.9), Inches(2.05), Inches(11.5), Inches(2.2))
    para(tf, "Satellite-Based Urban Growth and", 36, TEXT, True, 2, first=True)
    para(tf, "Economic Activity Intelligence System", 36, TEXT, True, 10)
    para(tf, "Review 3  -  Architecture, Methods and Results  -  Varanasi, Uttar Pradesh",
         16, ACCENT, False, 0)
    bar = s.shapes.add_shape(1, Inches(0.9), Inches(4.45), Inches(2.0), Inches(0.04))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()
    bar.shadow.inherit = False
    tf = tb(s, Inches(0.9), Inches(4.8), Inches(11.5), Inches(1.6))
    para(tf, "Mayank  23BCE1753        Achal Pramod Tripathi  23BCE1734        "
             "Mohammad Owais  23BCE1746", 14, TEXT, False, 6, first=True)
    para(tf, "Guide: Joshan        BCSE497J Project I  -  SCOPE  -  Fall 2026-27",
         13, DIM, False, 6)
    para(tf, "Panel review, 16 September 2026", 13, DIM, False, 0)
    notes(s, "Introduce the team and say who presents which part: architecture and "
             "pipeline, the typology and the growth model, validation and results.")

    # 2 - what the system answers ------------------------------------------
    s = blank(prs)
    heading(s, "What the system answers", "Three questions about Varanasi, from open satellite data")
    qs = [
        ("Where and how fast did the city grow?", ACCENT,
         "Built-up change from 2010 to 2020, measured on a 100 m grid, and the form it took: "
         "infill, edge expansion or leapfrog."),
        ("Is the new development being used?", WARN,
         "Night-time light, points of interest and population, compared with how much is "
         "built. New land that stays dark and empty is flagged as ghost growth."),
        ("Where will it grow next?", GREEN,
         "A growth model learns from 2010-2015, is tested on 2015-2020, and predicts new "
         "urban land for 2025 and 2030."),
    ]
    x = Inches(0.7)
    for q, colour, body in qs:
        c = card(s, x, Inches(1.75), Inches(3.85), Inches(2.5))
        c.line.color.rgb = BORDER
        tf = tb(s, x + Inches(0.22), Inches(1.92), Inches(3.45), Inches(2.25))
        para(tf, q, 16, colour, True, 8, first=True)
        para(tf, body, 13.5, DIM, False, 0)
        x += Inches(4.02)
    facts = [("1,206 km2", "study area around the city"), ("100 m", "analysis grid"),
             ("500 m", "reporting grid"), ("2010-2020", "observed built-up years"),
             ("2013-2024", "night-light years")]
    x = Inches(0.7)
    for v, lab in facts:
        chip(s, x, Inches(4.55), Inches(2.25), Inches(1.1), v, lab)
        x += Inches(2.41)
    footer(s, "Ghost growth: land that became built-up after 2010 but shows much less activity "
              "than comparable built-up land in the same city.")
    notes(s, "The problem: Varanasi is growing, but satellite built-up maps only say that land "
             "was built, not whether anyone uses it. We combine built-up area with signals of "
             "use - night-time light as a proxy for activity, points of interest, population - "
             "to separate new development that is filling up from development that is not. "
             "The third part predicts where growth goes next.")

    # 3 - what changed ------------------------------------------------------
    s = blank(prs)
    heading(s, "What changed since Review 2", "Three things, in order of importance")
    items = [
        ("1. We fixed what was wrong", ACCENT,
         "An audit of our own work found 15 defects. The most serious: GHSL's 2025 epoch - a "
         "projection, not an observation - was used as the current year, so every 2010-2025 "
         "figure was measured against a forecast. Every number is now recomputed on the "
         "observed epochs 2010, 2015 and 2020."),
        ("2. We tested the results against independent data", GREEN,
         "Six built-up datasets, a second temperature sensor, a hold-out test of the "
         "ghost-growth classes, and Indian records: Census of India 2001 and 2011, the 2013 "
         "Economic Census, and district GDP from the Government of Uttar Pradesh."),
        ("3. We finished the prediction module", WARN,
         "Logistic regression and random forest, trained on 2010-2015, scored on a held-out "
         "2015-2020; the better one predicts growth 5 and 10 years ahead."),
    ]
    y = Inches(1.75)
    for title, colour, body in items:
        c = card(s, Inches(0.7), y, Inches(11.9), Inches(1.55))
        c.line.color.rgb = BORDER
        tf = tb(s, Inches(0.95), y + Inches(0.16), Inches(11.4), Inches(1.25))
        para(tf, title, 16, colour, True, 4, first=True)
        para(tf, body, 13, DIM, False, 0)
        y += Inches(1.72)
    notes(s, "Lead with the correction: it is the strongest evidence of technical accuracy "
             "that we found and fixed our own mistake. Every number after this slide is on "
             "observed data only.")

    # 4 - implementation status ---------------------------------------------
    s = blank(prs)
    heading(s, "Implementation status", "What works today")
    rows = [
        ["Component", "Status", "What it delivers"],
        ["Data acquisition - GHSL, OpenStreetMap, 13 Earth Engine layers", "Done",
         "every layer on one 100 m grid"],
        ["Indian datasets - Census, Economic Census, UP district GDP", "New",
         "independent checks from Indian records"],
        ["Analysis - built-up, activity, green cover, heat, typology", "Done",
         "change maps and a six-class growth typology"],
        ["Growth models - logistic regression, random forest", "Extended",
         "held-out test; 2025 and 2030 predictions"],
        ["Validation - hold-out test, cross-checks, census, economy", "New",
         "six claims, each tested on independent data"],
        ["Results - 14 figures, 23 zone evidence cards", "New",
         "every number regenerated in one run"],
        ["Dashboard - seven views, including validation", "Updated", "interactive maps"],
        ["City search; image labels for flagged zones", "Review 4", "-"],
    ]
    grid_table(s, Inches(0.7), Inches(1.75), Inches(11.9), rows, [0.5, 0.12, 0.38])
    y = Inches(5.5)
    chip(s, Inches(0.7), y, Inches(3.8), Inches(1.1), "41 of 41",
         "automated tests pass; 16 added for Review 3")
    chip(s, Inches(4.75), y, Inches(3.8), Inches(1.1), "about 3.5 min",
         "one end-to-end run, tests to figures", GREEN)
    chip(s, Inches(8.8), y, Inches(3.8), Inches(1.1), "15 of 15",
         "audit defects fixed", WARN)
    notes(s, "Everything above the last row runs today and can be shown live. The whole chain "
             "- tests, analysis, models, validation, figures - runs as one step in about three "
             "and a half minutes on the data already downloaded.")

    # 5 - architecture ------------------------------------------------------
    diagram_slide(
        prs, "System architecture",
        "Five layers; data moves top to bottom and every arrow is a real hand-off",
        "A1_system_architecture.png",
        "Three kinds of source: open downloads with no account (GHSL built-up and population, "
        "OpenStreetMap), Google Earth Engine (night light, Sentinel-2, Landsat, MODIS, Dynamic "
        "World, Open Buildings, WorldCover, SRTM, WorldPop) and Indian records. The second "
        "layer puts every source on one 100 m grid; the 500 m reporting grid nests exactly, "
        "25 cells to one. Earth Engine computes on Google's servers and sends back only the "
        "clipped study area, which is why the whole system runs on a laptop.")

    # 6 - datasets ----------------------------------------------------------
    s = blank(prs)
    heading(s, "Datasets", "Official identifiers; the Indian records are new this review")
    sat = [
        ["Satellite and gridded dataset", "Used for"],
        ["GHS-BUILT-S R2023A, GHS-POP R2023A", "built-up surface, population, 100 m"],
        ["NOAA/VIIRS/DNB/ANNUAL_V21 + ANNUAL_V22", "night-time light, 2013-2024"],
        ["COPERNICUS/S2_SR_HARMONIZED", "NDVI, NDBI, true colour"],
        ["LANDSAT/LC08/C02/T1_L2 + LANDSAT/LC09/C02/T1_L2", "land surface temperature"],
        ["MODIS/061/MOD11A2", "independent temperature check"],
        ["GOOGLE/DYNAMICWORLD/V1", "water mask, built-up gain"],
        ["GOOGLE/Research/open-buildings-temporal/v1", "buildings in flagged cells"],
        ["ESA/WorldCover/v200", "independent built-up map"],
        ["USGS/SRTMGL1_003", "terrain slope"],
        ["WorldPop/GP/100m/pop", "independent population"],
    ]
    grid_table(s, Inches(0.7), Inches(1.75), Inches(7.0), sat, [0.62, 0.38], 10, 10.5)
    ind = [
        ["Indian and other records", "Used for"],
        ["Census of India 2001 and 2011", "population check, 71 districts"],
        ["Economic Census 2013", "non-farm jobs per town and village"],
        ["SHRUG, Development Data Lab", "town and village boundaries"],
        ["District Domestic Product, DES, Govt of UP", "district GDP 2020-21, 2021-22"],
        ["OpenStreetMap", "1,999 points of interest, 38,455 roads"],
        ["FAO/GAUL/2015/level2", "district boundaries"],
    ]
    grid_table(s, Inches(7.95), Inches(1.75), Inches(4.65), ind, [0.52, 0.48], 10, 10.5)
    footer(s, "SHRUG is used under CC BY-NC-SA 4.0 and cited as Asher, Lunt, Matsuura and "
              "Novosad (2021). DES: Directorate of Economics and Statistics.")
    notes(s, "Every dataset is named by its official identifier. Night-time light is used as "
             "an indirect indicator of activity - a proxy, not a measurement of the economy. "
             "The Indian records are used only to check our results, never as inputs.")

    # 7 - pipeline ----------------------------------------------------------
    diagram_slide(
        prs, "Processing flow", "Seven stages; each hands a named layer to the next",
        "A2_pipeline_dataflow.png",
        "Stage 1, built-up, feeds almost everything: the 2020 built fraction for vegetation, "
        "the 2010-2020 gain for the typology, and the mask of land already urban in 2010 that "
        "defines 'the established city' for the light trend. Stage 6 is where the three "
        "signals meet. Before anything runs, a check refuses any GHSL year after 2020 as a "
        "measurement - that is the Review 2 defect, fixed in the system itself.")

    # 8 - grid --------------------------------------------------------------
    diagram_slide(
        prs, "Analysis grid and aggregation",
        "Analysis at 100 m, reporting at 500 m - the scale night-time light can support",
        "A9_grid_and_aggregation.png",
        "VIIRS pixels are about 463 m, so 500 m is the honest reporting scale for activity, "
        "while 100 m keeps built-up change sharp. Totals are summed, averages averaged, yes/no "
        "findings use 'any', class maps use 'priority'. A majority vote erased every "
        "ghost-growth cell from the map, because a finding is always a minority of its block.")

    # 9 - temporal design ---------------------------------------------------
    diagram_slide(
        prs, "Temporal design",
        "2020 is the last observed epoch; later years are projections or our own predictions",
        "A8_temporal_design.png",
        "Built-up change is 2010 to 2020, observed only. The growth model trains on 2010-2015 "
        "and is tested on 2015-2020; the typology classifies on 2013-2020 and is checked on "
        "2021-2024. In both, the check uses years the method never saw. Activity uses the "
        "2022-2024 light mean, which gives land built by 2020 two years to be occupied.")

    # 10 - typology ---------------------------------------------------------
    diagram_slide(
        prs, "The ghost-growth algorithm",
        "Each cell is judged against comparably built land in the same city",
        "A3_ghost_typology_algorithm.png",
        "The activity index ranks each signal per square kilometre built - light 0.45, points "
        "of interest 0.35, population 0.20 - so a large built-up area is not 'active' just "
        "because it is large. Expected activity is the median of cells with a similar built-up "
        "fraction. New, under-used cells then split on the light trend: emerging if it rises "
        "significantly faster than the city, otherwise ghost growth. It is a screening rule; "
        "its precision is unmeasured until image labels exist.")

    # 11 - relative trend ---------------------------------------------------
    diagram_slide(
        prs, "Why the light trend is relative to the city",
        "Raw radiance rises everywhere; relative to the city, only some cells catch up",
        "A4_relative_light_trend.png",
        f"On raw radiance, {g('nightlights.urban_cells_significant_growth_pct', 51.4)}% of urban "
        "cells show a significant rise, so 'rising' on the raw sign says little - Review 2 used "
        "that rule. The 2021 to 2022 change of product version adds a step. Subtracting the "
        "established city's median cancels both, and leaves which cells are catching up.")

    # 12 - heat method ------------------------------------------------------
    diagram_slide(
        prs, "Surface heat-island method",
        "Land surface temperature minus a rural reference from the same scene",
        "A6_heat_island_method.png",
        "Cloud, shadow and cirrus are masked and every composite needs at least four dates. "
        "The rural reference now excludes the Ganga and recently built land; that moved it "
        "from 41.18 to 41.42 degrees C. The mean is taken over urban cells only.")

    # 13 - growth model -----------------------------------------------------
    diagram_slide(
        prs, "Growth-model architecture",
        "Learn where the city grew, test on years it never saw, then predict",
        "A5_growth_model.png",
        "Logistic regression is a straight-line model with readable weights; random forest "
        "combines 300 decision trees and captures non-linear effects. Both get the same eight "
        "drivers and the same split. A cellular automaton places the demanded number of cells, "
        "mixing suitability (0.65) with built-up nearby (0.35), so growth attaches to growth. "
        "We choose the model on Figure of Merit because it scores the placed map.")

    # 14 - validation framework ---------------------------------------------
    diagram_slide(
        prs, "Validation framework",
        "Every headline claim is checked against evidence the method never used",
        "A7_validation_framework.png",
        "Green means the independent evidence supports the claim; amber means it supports it "
        "with a stated qualification. No evidence here was used to build the result it checks.")

    # 15 - results: growth --------------------------------------------------
    s = blank(prs)
    heading(s, "Result: urban growth, 2010-2020", "Observed epochs only")
    place(s, FIG / "F01_growth_index.png", Inches(0.7), Inches(1.7), Inches(6.2), Inches(4.6))
    place(s, FIG / "F02_urban_form.png", Inches(7.2), Inches(1.7), Inches(5.4), Inches(2.7))
    xa, xb = Inches(7.2), Inches(9.95)
    chip(s, xa, Inches(4.55), Inches(2.6), Inches(1.0),
         f"+{b.get('net_builtup_change_km2', 17.25)} km2",
         f"built-up, {built.get('2010', 72.48)} to {built.get('2020', 89.73)} km2")
    chip(s, xb, Inches(4.55), Inches(2.65), Inches(1.0),
         f"{b.get('annual_urban_growth_pct', 1.86)} % / yr", "urban extent growth")
    chip(s, xa, Inches(5.7), Inches(2.6), Inches(1.0), "44.3 %",
         "of new urban land is leapfrog", WARN)
    chip(s, xb, Inches(5.7), Inches(2.65), Inches(1.0), "1.3x - 2.3x",
         "faster than population", GREEN)
    notes(s, "Built-up surface grew 23.8% in ten years while population grew 10 to 18% "
             "depending on the dataset, so land is consumed faster than people arrive: 1.3 to "
             "2.3 times, reported as a range because the two population datasets disagree. "
             "Almost half the new urban land is leapfrog - detached from the existing city. "
             "The 2025 values are drawn dashed: GHSL's projection, not data.")

    # 16 - results: ghost growth ------------------------------------------------
    s = blank(prs)
    heading(s, "Result: is the ghost-growth split real?",
            "Classified on 2013-2020, checked on 2021-2024 data the classifier never saw")
    place(s, FIG / "F04_rule_sensitivity.png", Inches(0.7), Inches(1.75), Inches(5.9), Inches(3.4))
    place(s, FIG / "F05_holdout.png", Inches(6.8), Inches(1.75), Inches(5.8), Inches(3.4))
    card(s, Inches(0.7), Inches(5.3), Inches(11.9), Inches(1.3))
    tf = tb(s, Inches(0.95), Inches(5.42), Inches(11.4), Inches(1.1))
    para(tf, "Yes - modestly. Cells called emerging went on to brighten more than cells called "
             "ghost growth.", 14, TEXT, True, 4, first=True)
    para(tf, f"p = {g('typology_holdout.cells.p_value', 0.0104):.3f} over 100 m cells and "
             f"p = {g('typology_holdout.blocks_500m.p_value', 0.0306):.3f} over 500 m blocks "
             "(one-sided Mann-Whitney). An emerging cell out-brightened a ghost cell 56% of the "
             "time: the split carries real information, and it is a screening map, not a verdict.",
         12.5, DIM, False, 0)
    notes(s, "Left: how the result depends on the definition of 'rising'. Our rule - faster than "
             "the city and significant - gives 4.12 square kilometres emerging and 1.44 ghost "
             "growth. Right: the hold-out test. The outcome was never seen by the classifier.")

    # 17 - results: a zone ----------------------------------------------------
    s = blank(prs)
    heading(s, "What a flagged zone looks like", "One evidence card of 23")
    place(s, FIG / "zones" / "zone_01.png", Inches(0.7), Inches(1.7), Inches(8.4), Inches(5.0))
    x = Inches(9.45)
    ob = R.get("open_buildings", {})
    share = g("open_buildings.classes.ghost_growth.share_with_buildings_2023", 0.819)
    chip(s, x, Inches(1.85), Inches(3.15), Inches(1.05), str(g("typology.n_zones", 23)),
         "ghost-growth zones flagged")
    chip(s, x, Inches(3.05), Inches(3.15), Inches(1.05), f"{share * 100:.0f} %",
         "of ghost cells contain buildings in 2023", GREEN)
    chip(s, x, Inches(4.25), Inches(3.15), Inches(1.05),
         f"{ob.get('ghost_cells_without_buildings_2023', 26)} of {ob.get('ghost_cells', 144)}",
         "ghost cells with no buildings: likely GHSL errors", WARN)
    footer(s, "Evidence for a reader, not a validation: how often a flagged zone is really vacant "
              "stays unmeasured until image labels exist (Review 4).")
    notes(s, "Each card shows Sentinel-2 true colour in 2018 and 2024, Open Buildings presence in "
             "2016 and 2023, and the zone's night-light series against the city. This zone was "
             "built up, but its light has stayed flat relative to the city since 2014.")

    # 18 - results: growth model --------------------------------------------
    s = blank(prs)
    heading(s, "Result: where the city grows next", "Trained 2010-2015, scored on a held-out 2015-2020")
    place(s, FIG / "F07_figure_of_merit.png", Inches(0.7), Inches(1.7), Inches(5.9), Inches(3.65))
    place(s, FIG / "F09_projection_map.png", Inches(6.9), Inches(1.7), Inches(3.7), Inches(3.55))
    rows = [
        ["Model", "AUC, held out", "Figure of Merit", "vs random"],
        ["Random forest", f"{rf.get('auc_test', 0.834):.3f}", f"{rf.get('figure_of_merit', 0.1031):.3f}",
         f"{rf.get('skill_vs_random', 18.72):.1f}x"],
        ["Logistic regression", f"{lr.get('auc_test', 0.9096):.3f}", f"{lr.get('figure_of_merit', 0.0687):.3f}",
         f"{lr.get('skill_vs_random', 12.46):.1f}x"],
        ["Random allocation", "0.500", f"{g('growth_models.random.figure_of_merit', 0.00551):.4f}", "1x"],
    ]
    grid_table(s, Inches(0.7), Inches(5.5), Inches(5.9), rows, [0.36, 0.2, 0.24, 0.2], 10, 10)
    chip(s, Inches(10.8), Inches(1.85), Inches(1.8), Inches(1.05), f"+{p25}", "km2 by 2025", GREEN)
    chip(s, Inches(10.8), Inches(3.05), Inches(1.8), Inches(1.05), f"+{p30}", "km2 by 2030", GREEN)
    tf = tb(s, Inches(6.9), Inches(5.45), Inches(5.7), Inches(1.4))
    para(tf, "Predictions, not observations.", 13, WARN, True, 3, first=True)
    para(tf, "The forest places growth better; the straight-line model ranks cells better "
             "overall. Removing roads changes nothing, so no later information leaks in. "
             "Published benchmark on the same measure: 0.264, PLUS model in Wuhan.",
         11.5, DIM, False, 0)
    notes(s, "Figure of Merit measures how much of the real 2015-2020 growth the placed map "
             "caught. Random forest catches 18.7 times more than random placement. AUC ranks "
             "all cells; FoM scores the map we actually publish, so we choose on FoM. The "
             "demand assumes the 2010-2020 growth rate of 1.86% a year continues.")

    # 19 - results: heat ------------------------------------------------------
    s = blank(prs)
    heading(s, "Result: surface heat", "March to May, daytime")
    place(s, FIG / "F12_land_surface_temperature.png", Inches(0.7), Inches(1.7), Inches(11.9), Inches(3.9))
    c = card(s, Inches(0.7), Inches(5.75), Inches(11.9), Inches(1.15))
    c.line.color.rgb = WARN
    tf = tb(s, Inches(0.95), Inches(5.85), Inches(11.4), Inches(1.0))
    para(tf, "Varanasi has no daytime surface heat island in the pre-monsoon season.",
         14, WARN, True, 3, first=True)
    para(tf, f"Urban cells are {abs(g('thermal.mean_urban_intensity_c', -1.66))} degrees C cooler "
             "than the bare farmland around them, MODIS agrees, and 22.3 of 23.8 km2 of hotspots "
             "lie on rural land.", 12.5, DIM, False, 0)
    notes(s, "Review 2 reported +1.10 degrees because it averaged only the cells warmer than "
             "rural - positive by construction. Dry, bare fields heat faster than the city by "
             "day before the monsoon; this pattern is reported for Indian cities in the "
             "literature (Shastri et al. 2017). Heat vulnerability, which weights heat by "
             "residents, remains the planning layer.")

    # 20 - results: Indian records ----------------------------------------------
    s = blank(prs)
    heading(s, "Checked against Indian records",
            "Census of India, Economic Census, Uttar Pradesh district GDP")
    place(s, FIG / "F10_population.png", Inches(0.7), Inches(1.72), Inches(5.9), Inches(2.6))
    place(s, FIG / "F13_lights_vs_economy.png", Inches(6.8), Inches(1.72), Inches(5.8), Inches(2.6))
    card(s, Inches(0.7), Inches(4.6), Inches(5.9), Inches(2.25))
    tf = tb(s, Inches(0.92), Inches(4.75), Inches(5.5), Inches(2.0))
    para(tf, "Population", 14, ACCENT, True, 4, first=True)
    bullets(tf, [
        "WorldPop is closer to the 2011 Census than GHS-POP: median district error -0.4% "
        "against +3.4%, over 71 districts.",
        "GHS-POP puts 63% too many people inside the city boundary.",
        "So growth is reported as a range: built-up grew 1.3x to 2.3x faster than population.",
    ], 11, 5)
    c = card(s, Inches(6.8), Inches(4.6), Inches(5.8), Inches(2.25))
    c.line.color.rgb = GREEN
    tf = tb(s, Inches(7.02), Inches(4.75), Inches(5.4), Inches(2.0))
    para(tf, "Night-time light as an activity proxy", 14, GREEN, True, 4, first=True)
    bullets(tf, [
        "District scale: light against UP district GDP, Spearman 0.85, R2 0.75.",
        "Village scale: light against 2013 Economic Census jobs, 0.41 - no better than "
        "population.",
        "That is why the screen reports zones, not single cells.",
    ], 11, 5)
    notes(s, "These are the Indian sources the panel asked about. They are used only as checks. "
             "The district result matches Henderson et al. (2012): light tracks economic output "
             "well at district scale.")

    # 21 - corrections --------------------------------------------------------
    diagram_slide(
        prs, "What the audit changed",
        "Eight numbers, Review 2 (hollow) against Review 3 (filled)",
        "P3_corrections_before_after.png",
        "New built-up surface fell from 23.47 to 17.25 square kilometres because Review 2 "
        "counted GHSL's 2025 projection as data. Heat intensity changed sign once the mean was "
        "taken over urban cells and water left the rural reference. AUC fell from 0.99 to 0.91 "
        "because 0.99 was measured on the training data.")

    # 22 - roadmap -------------------------------------------------------------
    diagram_slide(
        prs, "Where we are", "Six phases done; Review 4 adds city search, image labels and planning data",
        "P1_roadmap.png",
        "Everything left of the 'today' line is done. The hatched bars are planned, not claimed.")

    # 23 - limitations and next -------------------------------------------------
    s = blank(prs)
    heading(s, "Limitations, and what Review 4 adds", "Said plainly")
    c = card(s, Inches(0.7), Inches(1.72), Inches(5.85), Inches(4.95))
    c.line.color.rgb = WARN
    tf = tb(s, Inches(0.95), Inches(1.88), Inches(5.35), Inches(4.65))
    para(tf, "What we cannot claim yet", 15, WARN, True, 8, first=True)
    bullets(tf, [
        "No ground truth for the ghost flag: how often a flagged zone is really vacant is "
        "unmeasured. The hold-out test and the building check are indirect.",
        "VIIRS pixels are 463 m; at village scale, light mostly measures settlement size.",
        "GHS-POP over-counts the city core; the activity index still uses it, at weight 0.20.",
        "Heat is daytime only, and the 2013 composite has just 5 dates.",
        "The prediction assumes the 2010-2020 growth rate continues.",
    ], 11.5, 7)
    c = card(s, Inches(6.75), Inches(1.72), Inches(5.85), Inches(4.95))
    c.line.color.rgb = GREEN
    tf = tb(s, Inches(7.0), Inches(1.88), Inches(5.35), Inches(4.65))
    para(tf, "Review 4", 15, GREEN, True, 8, first=True)
    bullets(tf, [
        ("City search", "type any city name and run the same analysis"),
        ("Image labels", "read a sample of flagged and unflagged cells off historical "
                         "high-resolution imagery, and measure precision"),
        ("Planning data", "ward-level reporting, VDA Master Plan 2031, Bhuvan land use"),
        ("Model", "tune the neighbourhood weight; a range of demand scenarios"),
        ("Publication", "a paper built on the hold-out and proxy-scale findings"),
    ], 11.5, 8)
    notes(s, "Say the limitations before the panel does. The main one is that the ghost flag "
             "has no ground truth yet; Review 4 measures its precision with image labels.")

    # 24 - close ----------------------------------------------------------------
    s = blank(prs)
    heading(s, "Three things to take away")
    takeaways = [
        (ACCENT, f"Varanasi's built-up surface grew {g('population_validation.study_area_2010_2020.built_surface_growth_pct', 23.8)}% "
                 "in 2010-2020, 1.3x to 2.3x faster than its population, and 44% of new urban "
                 "land is leapfrog."),
        (WARN, "Most low-activity new development is filling up; "
               f"{g('typology.areas_km2.ghost_growth', 1.44)} km2 in {g('typology.n_zones', 23)} "
               "zones is not - and that split holds up on data the method never saw."),
        (GREEN, f"A random forest places future growth best: +{p25} km2 by 2025 and +{p30} km2 "
                "by 2030 - predictions, not observations."),
    ]
    y = Inches(1.8)
    for colour, text in takeaways:
        c = card(s, Inches(0.7), y, Inches(11.9), Inches(1.2))
        c.line.color.rgb = BORDER
        tf = tb(s, Inches(0.95), y + Inches(0.2), Inches(11.4), Inches(0.9))
        para(tf, text, 15, colour, True, 0, first=True)
        y += Inches(1.38)
    tf = tb(s, Inches(0.7), Inches(6.0), Inches(11.9), Inches(0.9))
    para(tf, "Thank you. Questions are welcome.", 16, TEXT, True, 4, first=True)
    para(tf, "AI assistance (Claude, by Anthropic) was used in this project and is acknowledged "
             "as the project guidelines require; the team checked and presents every part.",
         10.5, DIM, False, 0)
    notes(s, "Questions to expect: why the 2025 epoch mattered; why the heat result reversed; "
             "why the forest wins on Figure of Merit but loses on AUC; what the ghost flag can "
             "and cannot claim.")

    tidy(prs)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    return OUT, len(prs.slides)


if __name__ == "__main__":
    path, n = build()
    print(f"wrote {path}  ({n} slides, {path.stat().st_size // 1024} KB)")
